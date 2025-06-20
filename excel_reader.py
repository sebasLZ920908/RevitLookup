# -*- coding: utf-8 -*-
"""
Módulo para leer datos de espacios desde un archivo Excel.
"""

import openpyxl

def leer_datos_espacios_excel(ruta_archivo_excel: str) -> list[dict]:
    """
    Lee los datos de espacios desde un archivo Excel.

    Args:
        ruta_archivo_excel: La ruta al archivo Excel (.xlsx).

    Returns:
        Una lista de diccionarios, donde cada diccionario representa un espacio
        y contiene los datos de ese espacio con claves correspondientes a los
        encabezados del Excel. Devuelve una lista vacía si ocurre un error.
    """
    datos_espacios = []
    try:
        # Abrir el libro de trabajo y seleccionar la hoja activa
        workbook = openpyxl.load_workbook(ruta_archivo_excel)
        sheet = workbook.active

        # Leer los encabezados de la primera fila
        headers_row = sheet[1]
        headers = [cell.value for cell in headers_row]

        # Validar que los encabezados esperados estén presentes
        expected_headers = ["Código del espacio", "Nombre del ambiente", "Área", "Longitud", "Ancho", "Altura"]
        if not all(header in headers for header in expected_headers):
            print(f"Error: El archivo Excel no contiene los encabezados esperados. Encabezados encontrados: {headers}")
            return []

        # Mapear los nombres de encabezado a los índices de columna
        header_to_index = {header: i for i, header in enumerate(headers)}

        # Columnas que necesitan conversión a float
        columnas_numericas = ["Área", "Longitud", "Ancho", "Altura"]

        # Iterar sobre las filas de datos (a partir de la segunda fila)
        for row_num in range(2, sheet.max_row + 1):
            row_data = {}
            try:
                current_row_values = sheet[row_num]
                for header in headers:
                    cell_value = current_row_values[header_to_index[header]].value
                    if header in columnas_numericas:
                        if cell_value is not None:
                            try:
                                row_data[header] = float(cell_value)
                            except ValueError:
                                print(f"Advertencia: No se pudo convertir '{cell_value}' a número para la columna '{header}' en la fila {row_num}. Se usará None.")
                                row_data[header] = None
                        else:
                            row_data[header] = None
                    else:
                        row_data[header] = cell_value
                datos_espacios.append(row_data)
            except Exception as e:
                print(f"Error procesando la fila {row_num}: {e}")
                # Se podría decidir continuar con otras filas o parar.
                # Por ahora, se omite la fila con error y se continúa.
                continue

    except FileNotFoundError:
        print(f"Error: No se encontró el archivo en la ruta especificada: {ruta_archivo_excel}")
        return []
    except Exception as e:
        print(f"Ocurrió un error general al procesar el archivo Excel: {e}")
        return []

    return datos_espacios

if __name__ == '__main__':
    # Crear un archivo Excel de ejemplo para pruebas
    # Esto normalmente no estaría en el mismo script, pero es útil para demostración.
    from openpyxl import Workbook

    # Crear un nuevo libro de trabajo y seleccionar la hoja activa
    wb_test = Workbook()
    ws_test = wb_test.active
    ws_test.title = "DatosEspacios"

    # Escribir los encabezados
    encabezados_test = ["Código del espacio", "Nombre del ambiente", "Área", "Longitud", "Ancho", "Altura", "Otra Columna"]
    ws_test.append(encabezados_test)

    # Escribir datos de ejemplo
    datos_ejemplo = [
        ("ESP001", "Oficina Principal", 50, 10, 5, 3.0, "Dato extra 1"),
        ("ESP002", "Sala de Reuniones", 25, 5, 5, 3, "Dato extra 2"),
        ("ESP003", "Recepción", 30.5, 6.0, 5.08, 2.8, "Dato extra 3"),
        ("ESP004", "Baño", "N/A", 2, 2.5, 2.8, "Dato extra 4"), # Área con error de tipo
        ("ESP005", "Cocina", 15, 4, None, 2.8, "Dato extra 5"), # Ancho con None
    ]

    for fila_dato in datos_ejemplo:
        ws_test.append(fila_dato)

    ruta_archivo_test = "espacios_test.xlsx"
    wb_test.save(ruta_archivo_test)
    print(f"Archivo Excel de prueba '{ruta_archivo_test}' creado.")

    # Probar la función
    print(f"\nLeyendo datos del archivo: {ruta_archivo_test}")
    datos_leidos = leer_datos_espacios_excel(ruta_archivo_test)

    if datos_leidos:
        print("\nDatos leídos exitosamente:")
        for espacio in datos_leidos:
            print(espacio)
    else:
        print("\nNo se pudieron leer datos o el archivo estaba vacío/con errores críticos.")

    # Probar con un archivo inexistente
    print(f"\nIntentando leer un archivo inexistente:")
    datos_inexistentes = leer_datos_espacios_excel("archivo_inexistente.xlsx")
    if not datos_inexistentes:
        print("Prueba de archivo inexistente completada como se esperaba.")

    # Limpiar el archivo de prueba (opcional)
    import os
    # os.remove(ruta_archivo_test)
    # print(f"\nArchivo de prueba '{ruta_archivo_test}' eliminado.")
